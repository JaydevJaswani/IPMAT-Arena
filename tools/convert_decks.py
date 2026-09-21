#!/usr/bin/env python3
"""
convert_decks.py — turn JJ's IPMAT practice decks into a structured static question bank.
Handles two clean formats:
  1) Numbers 1 Workshop  (Q slide -> Solution slide with ANSWER + Concept/Trap)
  2) Linear/Surds/Indices (QUESTION N OF 7 -> QUESTION N · SOLUTION with ANSWER)
Output: data/bank_static.json  +  a console report flagging anything to eyeball.
"""
import re, json, sys
from pathlib import Path
from pptx import Presentation

OUT = Path(__file__).resolve().parent.parent / "data"
OUT.mkdir(exist_ok=True)

def shapes(slide):
    return [sh.text_frame.text.strip() for sh in slide.shapes
            if sh.has_text_frame and sh.text_frame.text.strip()]

def is_footer(t): return "IMS Ahmedabad" in t or "Quants Portal" in t
def classify_topic(stem):
    s = stem.lower()
    if "√" in stem or "index" in s or "indices" in s or "ⁿ" in stem or re.search(r"2[²³ˣⁿ]", stem):
        return "Surds & Indices"
    if "remainder" in s or "mod " in s or "divided by" in s and "^" in stem:
        return "Number Theory"
    if any(w in s for w in ["age", "equation", "x +", "solve", "digits", "two-digit"]):
        return "Linear Equations" if any(w in s for w in ["age","equation","solve","x +","3x","2y"]) else "Number System"
    if "divisible" in s or "hcf" in s or "lcm" in s or "divides" in s: return "Number System"
    return "Number System"

def normalize_answer(raw):
    a = raw.replace("ANSWER", "", 1).strip()
    a = re.sub(r"^[:=\s]+", "", a)
    # pull "x = 5" / "= 46 years" style
    m = re.search(r"=\s*([^=]+?)\s*(years|$)", a)
    core = a
    if "=" in a:
        core = a.split("=")[-1].strip()
    core = re.sub(r"\b(years|year)\b", "", core).strip()
    numeric = bool(re.fullmatch(r"-?\d+(\.\d+)?", core))
    return {"display": a, "value": core, "numeric": numeric}

def parse_numbers1(path):
    p = Presentation(path); out = []; qslides = {}; sslides = {}
    for s in p.slides:
        sh = shapes(s)
        qn = next((int(m.group(1)) for t in sh for m in [re.fullmatch(r"Q(\d+)", t)] if m), None)
        if qn is None: continue
        if any(t == "Space for working" for t in sh): qslides[qn] = sh
        elif any(t.startswith("ANSWER") for t in sh): sslides[qn] = sh
    for qn in sorted(qslides):
        if qn not in sslides: continue
        q, s = qslides[qn], sslides[qn]
        stem = max([t for t in q if not re.fullmatch(r"Q\d+", t) and t != "Space for working" and not is_footer(t)], key=len, default="")
        ans = next((t for t in s if t.startswith("ANSWER")), "")
        trap = next((t for t in s if t.startswith("Concept / Trap")), "")
        sol = max([t for t in s if not re.fullmatch(r"Q\d+", t) and t not in ("Solution",) and not t.startswith("ANSWER") and not t.startswith("Concept / Trap") and not is_footer(t) and not t.endswith("· Solution")], key=len, default="")
        na = normalize_answer(ans)
        out.append({"source":"Numbers 1 Workshop","topic":classify_topic(stem),"tier":"Heavy & Tricky",
                    "type":"int" if na["numeric"] else "short","stem":stem,"options":None,
                    "answer":na["value"],"answer_display":na["display"],
                    "solution":sol,"trap":trap.replace("Concept / Trap:","").strip()})
    return out

def parse_linear(path):
    p = Presentation(path); out=[]; tier="Warm-Up"; qslides={}; sslides={}
    TIERMAP={"WARM-UP":"Warm-Up","EXAM-RELEVANT":"Exam-Relevant","HEAVY":"Heavy & Tricky"}
    order=[]  # (kind, num, tier)
    for s in p.slides:
        sh=shapes(s)
        for t in sh:
            m=re.search(r"SET\s*\d+\s*·?\s*(WARM-UP|EXAM-RELEVANT|HEAVY)", t.upper())
            if m: tier=TIERMAP[m.group(1)]
        mo=next((re.search(r"QUESTION\s+(\d+)\s+OF\s+\d+",t) for t in sh if re.search(r"QUESTION\s+\d+\s+OF",t)),None)
        so=next((re.search(r"QUESTION\s+(\d+)\s*·\s*SOLUTION",t) for t in sh if "SOLUTION" in t.upper() and "QUESTION" in t.upper()),None)
        if mo:
            n=int(mo.group(1))
            stem=""
            if "Question" in sh:
                i=sh.index("Question")
                for t in sh[i+1:]:
                    if t!="Your working" and not re.fullmatch(r"\d+",t): stem=t; break
            qslides[(tier,n)]={"stem":stem}
        if so:
            n=int(so.group(1))
            ans=next((t for t in sh if t.startswith("ANSWER")),"")
            qn_stem=next((t for t in sh if re.match(rf"Q{n}\.",t)),"")
            sol=max([t for t in sh if not t.startswith("ANSWER") and not t.startswith("SET") and "SOLUTION" not in t.upper() and not re.match(r"Q\d+\.",t)],key=len,default="")
            sslides[(tier,n)]={"answer":ans,"stem":re.sub(rf"^Q{n}\.\s*","",qn_stem),"sol":sol}
    for key in sorted(set(list(qslides)+list(sslides))):
        q=qslides.get(key,{}); s=sslides.get(key,{})
        stem=q.get("stem") or s.get("stem","")
        na=normalize_answer(s.get("answer",""))
        out.append({"source":"Linear/Surds/Indices","topic":classify_topic(stem),"tier":key[0],
                    "type":"int" if na["numeric"] else "short","stem":stem,"options":None,
                    "answer":na["value"],"answer_display":na["display"],
                    "solution":s.get("sol",""),"trap":""})
    return out

def main():
    decks=[("/Users/jj/Downloads/IMS_Numbers1_Workshop.pptx",parse_numbers1),
           ("/Users/jj/Downloads/IPMAT_Linear_Surds_Indices_Practice.pptx",parse_linear)]
    bank=[]
    for path,fn in decks:
        got=fn(path); bank.extend(got); print(f"{Path(path).name}: {len(got)} questions")
    OPEN_KW=["justify","counterexample","counter-example","true or false","prove","is this possible",
             "is it possible","explain why","give an example","give a ","how many ways","in how many ways can you"]
    def mode_of(q):
        s=q["stem"].lower(); a=q["answer"]
        if q["type"]=="int": return "auto"
        if ";" in a or a.lower() in ("false","true") or a.startswith(("no","yes","e.g")) or "impossible" in a.lower(): return "open"
        if any(k in s for k in OPEN_KW): return "open"
        if re.fullmatch(r"[0-9√/()+\-·^².\s]+", a): return "auto"   # clean math token: 7/8, 2√2, √5
        return "open"
    for i,q in enumerate(bank,1):
        q["id"]=f"S{i:03d}"; q["mode"]=mode_of(q)
    (OUT/"bank_static.json").write_text(json.dumps(bank,indent=2,ensure_ascii=False))
    # report
    flags=[q for q in bank if not q["stem"] or not q["answer"]]
    print(f"\nTOTAL static questions: {len(bank)}")
    from collections import Counter
    print("by tier:",dict(Counter(q['tier'] for q in bank)))
    print("by topic:",dict(Counter(q['topic'] for q in bank)))
    print("by type:",dict(Counter(q['type'] for q in bank)))
    print(f"NEEDS REVIEW (missing stem/answer): {len(flags)}")
    for q in flags: print("   ",q["id"],repr(q["stem"][:40]),"ans=",repr(q["answer"]))
    print("\nSample:")
    for q in bank[:3]: print("  ",q["id"],q["tier"],"|",q["stem"][:60],"=>",q["answer_display"][:30])
    print("wrote data/bank_static.json")

if __name__=="__main__": main()
