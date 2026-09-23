#!/usr/bin/env python3
"""
harvest_graded.py — parse JJ's "ENHANCED / GRADED PRACTICE" deck template.
Each practice slide: difficulty (EASY/MODERATE/ADVANCED) + numbered stem + a/b/c/d
options + SOLUTION + 'Answer: [x] value'. Also harvests RAPID FIRE recall slides.
Difficulty maps to levels: Easy→L1, Moderate→L2, Advanced→L3, PYQ→L4.
"""
import re, sys, json, argparse
from pptx import Presentation

DIFF = {"EASY": ("L1", "Warm-Up"), "MODERATE": ("L2", "Exam-Relevant"),
        "ADVANCED": ("L3", "Heavy & Tricky"), "PYQ": ("L4", "Killer"), "IPMAT": ("L4", "Killer")}

def slide_lines(s):
    out = []
    for sh in s.shapes:
        if sh.has_text_frame:
            for ln in sh.text_frame.text.split("\n"):
                if ln.strip():
                    out.append(ln.strip())
    return out

def harvest(path, topic):
    p = Presentation(path)
    bank = []
    for s in p.slides:
        L = slide_lines(s)
        blob = "\n".join(L)
        am = re.search(r"Answer\s*:?\s*\[?\s*([a-dA-D])\s*\]?\s*(.*)", blob)
        # ---- graded MCQ slide ----
        if am and re.search(r"\b(EASY|MODERATE|ADVANCED)\b", blob) and any(re.fullmatch(r"[a-d]", x) for x in L):
            diff = next((d for d in ("ADVANCED","MODERATE","EASY") if d in blob), "MODERATE")
            level, tier = DIFF[diff]
            # options: a letter line followed by its value line
            opts = {}
            for i, ln in enumerate(L):
                if re.fullmatch(r"[a-d]", ln) and i+1 < len(L):
                    opts[ln] = L[i+1].strip()
            if len(opts) < 3:
                continue
            # stem: the line(s) between the lone number and the first option 'a'
            stem = ""
            try:
                ai = next(i for i, ln in enumerate(L) if re.fullmatch(r"a", ln))
                # walk back from a to the numbered marker
                num_i = next((i for i in range(ai-1, -1, -1) if re.fullmatch(r"\d{1,2}", L[i])), None)
                if num_i is not None:
                    stem = " ".join(L[num_i+1:ai]).strip()
            except StopIteration:
                pass
            letter = am.group(1).lower()
            val = am.group(2).strip() or opts.get(letter, "")
            if not stem or letter not in opts or len(stem) < 10:
                continue
            sol = ""
            if "SOLUTION" in blob:
                sol = re.sub(r"\s+", " ", blob.split("SOLUTION", 1)[1].split("Answer")[0]).strip()[:400]
            bank.append({"topic": topic, "level": level, "tier": tier, "type": "mcq", "mode": "auto",
                         "stem": re.sub(r"\s+", " ", stem), "options": [opts[k] for k in sorted(opts)],
                         "answer": opts[letter], "answer_display": opts[letter], "solution": sol,
                         "source": path.split("/")[-1]})
        # ---- rapid fire recall slide: 'N \n stem \n Answer  value' ----
        elif "RAPID FIRE" in blob and blob.count("Answer") >= 3:
            for m in re.finditer(r"(?:^|\n)(\d{1,2})\n(.+?)\nAnswer\s+(.+?)(?=\n\d{1,2}\n|\Z)", blob, re.S):
                stem = re.sub(r"\s+", " ", m.group(2)).strip(); val = re.sub(r"\s+", " ", m.group(3)).strip()
                if len(stem) < 8 or len(val) > 30:
                    continue
                numeric = bool(re.fullmatch(r"-?\d+", val))  # TITA only for whole numbers
                bank.append({"topic": topic, "level": "L1", "tier": "Warm-Up",
                             "type": "int" if numeric else "short", "mode": "auto" if numeric else "open",
                             "stem": stem, "options": None, "answer": val, "answer_display": val,
                             "solution": "", "source": path.split("/")[-1]})
    # dedupe
    seen = set(); uniq = []
    for q in bank:
        k = q["stem"][:50].lower()
        if k not in seen:
            seen.add(k); uniq.append(q)
    return uniq

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx"); ap.add_argument("--topic", required=True); ap.add_argument("--out", default="")
    a = ap.parse_args()
    b = harvest(a.pptx, a.topic)
    from collections import Counter
    print(f"{a.pptx.split('/')[-1]}: {len(b)} Q  | levels {dict(Counter(q['level'] for q in b))} | modes {dict(Counter(q['mode'] for q in b))}")
    for q in b[:3]:
        print("  •", q["stem"][:66], "| opts", q["options"], "=>", q["answer"])
    if a.out: json.dump(b, open(a.out, "w"), indent=1, ensure_ascii=False)
