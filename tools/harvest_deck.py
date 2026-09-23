#!/usr/bin/env python3
"""
harvest_deck.py — pull Q + answer pairs from JJ's PPTX practice decks.
Heuristic: a slide carrying 'ANSWER'/'Ans:' is a solution slide; its stem is either the
'Qn. ...' line on it or the longest text on the nearest preceding question slide.
Prints yield; --out saves JSON. Tune per deck as needed.
"""
import re, sys, json, argparse
from pptx import Presentation

def slide_texts(s):
    return [sh.text_frame.text.strip() for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]

def is_footer(t):
    return bool(re.search(r"IMS Ahmedabad|Trusted for Success|Prepared by|Page \d|Jaydev|·  Q\d", t))

ANS = re.compile(r"(?:ANSWER|Ans\.?|Answer)\s*[:\-]?\s*(.+)", re.I)

def harvest(path, topic, level):
    p = Presentation(path)
    slides = [slide_texts(s) for s in p.slides]
    L2T = {"L1":"Warm-Up","L2":"Exam-Relevant","L3":"Heavy & Tricky","L4":"Killer"}
    out = []
    for i, sh in enumerate(slides):
        joined = "\n".join(sh)
        m = ANS.search(joined)
        if not m:
            continue
        ans = re.sub(r"\s+", " ", m.group(1)).strip()
        # strip prose like "x = 5 years" -> keep the value part but store display too
        # stem: prefer a "Qn. ..." line on this slide
        stem = ""
        for t in sh:
            mm = re.match(r"Q\s*\d+[.\)]?\s*(.+)", t, re.S)
            if mm and len(mm.group(1)) > 15:
                stem = mm.group(1); break
        if not stem:
            # look back up to 2 slides for a question stem
            for j in range(i-1, max(-1, i-3), -1):
                cand = [t for t in slides[j] if not is_footer(t) and "?" in t or (t and len(t) > 40)]
                cand = [t for t in slides[j] if not is_footer(t) and len(t) > 25 and not ANS.search(t)]
                if cand:
                    stem = max(cand, key=len); break
        stem = re.sub(r"\s+", " ", stem).strip()
        if not stem or len(stem) < 12 or not ans or len(ans) > 220:
            continue
        # solution = other long text on the slide that isn't the answer/stem
        sol = ""
        for t in sh:
            if not ANS.search(t) and not is_footer(t) and t[:20] not in stem[:20] and len(t) > 30:
                sol = re.sub(r"\s+", " ", t); break
        val = ans.split("=")[-1].strip() if "=" in ans else ans
        val = re.sub(r"\b(years?|cm|m|kmph|km/h|units?|Rs\.?)\b", "", val).strip()
        numeric = bool(re.fullmatch(r"-?\d+", val))
        out.append({"topic": topic, "level": level, "tier": L2T.get(level, "Exam-Relevant"),
                    "type": "int" if numeric else "short", "mode": "auto" if numeric or re.fullmatch(r"[0-9√/().+\-·^² ]+", val) else "open",
                    "stem": stem, "options": None, "answer": val, "answer_display": ans,
                    "solution": sol, "source": path.split("/")[-1]})
    # de-dupe by stem
    seen = {}; uniq = []
    for q in out:
        k = q["stem"][:60].lower()
        if k in seen: continue
        seen[k] = 1; uniq.append(q)
    return uniq

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx"); ap.add_argument("--topic", required=True); ap.add_argument("--level", default="L2"); ap.add_argument("--out", default="")
    a = ap.parse_args()
    b = harvest(a.pptx, a.topic, a.level)
    print(f"{a.pptx.split('/')[-1]}: {len(b)} Q harvested ({sum(1 for q in b if q['mode']=='auto')} auto)")
    for q in b[:3]: print("  •", q["stem"][:64], "=>", q["answer_display"][:24])
    if a.out: json.dump(b, open(a.out, "w"), indent=1, ensure_ascii=False)
